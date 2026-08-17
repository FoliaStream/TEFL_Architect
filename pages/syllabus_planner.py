import streamlit as st
import os 
import yaml 

from src.fe.support_functions import setup_sidebar
from src.fe.styles import HIDE_SIDEBAR_NAV, TEXT_JUSTIFIED
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from io import BytesIO


################
# --- SET UP ---
################

# --- PAGE CONFIG --- 
st.set_page_config(page_title="Syllabus Planner", layout="wide")

# --- STYLES ---
st.markdown(HIDE_SIDEBAR_NAV, unsafe_allow_html=True)
st.markdown(TEXT_JUSTIFIED, unsafe_allow_html=True)

# --- CONFIG ---
with open(f"{os.getcwd()}/src/be/config.yaml", "r") as config_file:
    config = yaml.safe_load(config_file)

# --- SIDEBAR & TITLE ---
selected_page = setup_sidebar(
    pages=config['pages'],
    main_page=config['main_page']
)

# Navigation on click
if selected_page == "HOME":
    st.switch_page("HOME.py")
elif selected_page == "Whiteboard":
    st.switch_page("pages/whiteboard.py")
elif selected_page == "IELTS Preparation":
    st.switch_page("pages/ielts_preparation.py")
elif selected_page == "Lesson Planner":
    st.switch_page("pages/lesson_planner.py")
elif selected_page == "Syllabus Planner":
    st.title(selected_page)
elif selected_page == "TEFL Theory":
    st.switch_page("pages/tefl_theory.py")
elif selected_page == "Discover":
    st.switch_page("pages/discover.py")



# --- SESSION STATE INIT ---

# Basic info - start with empty course plan
if "course_plan" not in st.session_state:
    st.session_state.course_plan = {
        "metadata":{
            "class":"",
            "teacher":"",
            "teacher_contact":"",
            "number_students":"",
            "level_students":"",
            "start_date":"today",
            "end_date":None,
            "title":""
        },
        "sections": []  # Start with empty sections
    }


# Templates - these are the STRUCTURES to load, not the actual data
if "syllabus_template" not in st.session_state:
    st.session_state.syllabus_template = {
        "Standard":[
            {"type":"text", "title":"Information", "value":""},
            {"type":"course_structure_table","title":"Course Structure", "units": []},
            {"type":"lesson_structure_table","title":"Generic Lesson Structure", "stages": []},
            {"type":"schedule", "title":"Schedule & Lesson Information", "lessons": []},
            {"type":"text", "title":"Homework", "value":""}
        ]
    }

# --- PAGE FUNCTIONS ---
def format_section_title(doc, title):

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add blue background
    p_pr = title_para._element.get_or_add_pPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), '4472C4')  # Blue
    p_pr.append(shading)
    
    # Add white text
    title_run = title_para.add_run(title)
    title_run.font.size = Pt(14)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(255, 255, 255)
    
    return title_para

def format_table(table, column_align = None, header_row = True):

    tbl = table._tbl
    tblPr = tbl.tblPr

    # Borders
    tblBorders = OxmlElement('w:tblBorders')

    border_types = {
        'top': ('single', 12),
        'bottom': ('single', 8),
        'left': ('nil', 0),
        'right': ('nil', 0),
        'insideH': ('nil', 0),
        'insideV': ('nil', 0)
    }

    for border_name, (val, sz) in border_types.items():
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        if val != 'nil':
            border.set(qn('w:sz'), str(sz))
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), '000000') # black
        tblBorders.append(border)

    existing_borders = tblPr.find(qn('w:tblBorders'))
    if existing_borders is not None:
        tblPr.remove(existing_borders)
    tblPr.append(tblBorders)

    for row in table.rows:
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcBorders = tcPr.find(qn('w:tcBorders'))
            if tcBorders is not None:
                tcPr.remove(tcBorders)

    # Text Format
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):

            if column_align and col_idx < len(column_align):
                h_align, v_align = column_align[col_idx]
            else:
                #default
                h_align, v_align = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_VERTICAL.TOP
            
            # Set alignment 
            cell.vertical_alignment = v_align
            for paragraph in cell.paragraphs:
                paragraph.alignment = h_align

            
            # Set header bold
            if header_row and row_idx == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True

    return table


def format_syllabus_document(data):
    # Create doc
    doc = Document()

    # Set margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # ==== Page 1 ====
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    course_title = data['metadata']['title']
    header_run = header.add_run(f"{course_title}\nCOURSE PLAN")
    header_run.font.size = Pt(16)
    header_run.font.bold = True
    header_run.font.color.rgb = RGBColor(0, 51, 102)

    # Whitespace
    doc.add_paragraph() 

    # Process each section
    for section in data['sections']:
        # Add section title
        format_section_title(doc, section['title'])
        doc.add_paragraph()

        if section['type'] == 'text':
            # Simple text content
            content = section.get('value', '')
            if isinstance(content, list):
                for item in content:
                    if item:
                        doc.add_paragraph(item)
            elif content:
                doc.add_paragraph(content)
            doc.add_paragraph()

        elif section['type'] == 'course_structure_table':
            # Course structure table with units and lessons
            structure_table = doc.add_table(rows=1, cols=3)
            structure_table.style = "Table Grid"
            
            structure_table = format_table(structure_table,
                                          column_align=[(WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_VERTICAL.CENTER),
                                                       (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_VERTICAL.CENTER),
                                                       (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_VERTICAL.TOP)],
                                          header_row=True)
            
            headers = structure_table.rows[0].cells
            headers[0].text = "UNIT"
            headers[1].text = "LESSON TITLE"
            headers[2].text = "LESSON FOCUS"

            units = section.get('units', [])
            for unit in units:
                if unit.get('lessons'):
                    # Add unit header
                    unit_cell = structure_table.add_row().cells
                    unit_cell[0].text = f"Unit {unit.get('number', '')}\n{unit.get('title', '')}"
                    unit_cell[0].paragraphs[0].runs[0].font.bold = True
                    unit_cell[1].text = ""
                    unit_cell[2].text = unit.get('description', '')
                    
                    # Add lessons
                    for lesson in unit.get('lessons', []):
                        row = structure_table.add_row().cells
                        row[0].text = f"{unit.get('number', '')}.{lesson.get('number', '')}"
                        row[1].text = lesson.get('title', '')
                        row[2].text = lesson.get('focus', '')
            
            doc.add_paragraph()

        elif section['type'] == 'lesson_structure_table':
            # Generic lesson structure table
            structure_table = doc.add_table(rows=1, cols=2)
            structure_table.style = "Table Grid"
            
            structure_table = format_table(structure_table,
                                          column_align=[(WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_VERTICAL.CENTER),
                                                       (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_VERTICAL.TOP)],
                                          header_row=True)
            
            headers = structure_table.rows[0].cells
            headers[0].text = "STAGE"
            headers[1].text = "DESCRIPTION"

            stages = section.get('stages', [])
            for stage in stages:
                row = structure_table.add_row().cells
                row[0].text = stage.get('name', '')
                row[0].paragraphs[0].runs[0].font.bold = True
                row[1].text = stage.get('description', '')
            
            doc.add_paragraph()

        elif section['type'] == 'schedule':
            # Schedule and lesson information
            lessons = section.get('lessons', [])
            for lesson in lessons:
                # Create a bordered box for each lesson
                lesson_table = doc.add_table(rows=1, cols=1)
                lesson_table.style = "Table Grid"
                
                lesson_cell = lesson_table.cell(0, 0)
                
                # Format lesson content
                content = []
                if lesson.get('code'):
                    content.append(f"Lesson {lesson['code']}")
                if lesson.get('date'):
                    content.append(f"Date: {lesson['date']}")
                if lesson.get('time'):
                    content.append(f"Time: {lesson['time']}")
                if lesson.get('title'):
                    content.append(f"Title: {lesson['title']}")
                if lesson.get('topic'):
                    content.append(f"Topic: {lesson['topic']}")
                if lesson.get('description'):
                    content.append(f"Description: {lesson['description']}")
                if lesson.get('objectives'):
                    content.append(f"Learning Objectives:\n{lesson['objectives']}")
                if lesson.get('activity'):
                    content.append(f"Interactive activity: {lesson['activity']}")
                
                lesson_cell.text = "\n".join(content)
                doc.add_paragraph()

    return doc


# --- PAGE UI ---

# Input Basic Information
st.divider()

col1, col2 = st.columns(2)
with col1:
    st.markdown("### Course Information")
    st.session_state.course_plan["metadata"]["title"] = st.text_input(
        "Course Title", 
        value=st.session_state.course_plan["metadata"]["title"],
        key="syllabus_course_title"
    )
    st.session_state.course_plan["metadata"]["class"] = st.text_input(
        "Class/Level", 
        value=st.session_state.course_plan["metadata"]["class"],
        key="syllabus_class"
    )
    st.session_state.course_plan["metadata"]["number_students"] = st.text_input(
        "Number of Students", 
        value=st.session_state.course_plan["metadata"]["number_students"],
        key="syllabus_num_students"
    )
    st.session_state.course_plan["metadata"]["level_students"] = st.text_input(
        "Level of Students", 
        value=st.session_state.course_plan["metadata"]["level_students"],
        key="syllabus_level"
    )

with col2:
    st.markdown("### Schedule & Teacher")
    st.session_state.course_plan["metadata"]["start_date"] = st.date_input(
        "Start Date", 
        value=st.session_state.course_plan["metadata"]["start_date"],
        key="syllabus_start_date"
    )
    st.session_state.course_plan["metadata"]["end_date"] = st.date_input(
        "End Date", 
        value=st.session_state.course_plan["metadata"]["end_date"],
        key="syllabus_end_date"
    )
    st.session_state.course_plan["metadata"]["teacher"] = st.text_input(
        "Teacher Name", 
        value=st.session_state.course_plan["metadata"]["teacher"],
        key="syllabus_teacher"
    )
    st.session_state.course_plan["metadata"]["teacher_contact"] = st.text_input(
        "Teacher Contact", 
        value=st.session_state.course_plan["metadata"]["teacher_contact"],
        key="syllabus_contact"
    )

st.divider()


# Template loading section
col1, col2 = st.columns([3,1])
with col1: 
    template_choice = st.selectbox(
        "Load Template (this will replace your current sections)",
        options=list(st.session_state.syllabus_template.keys()),
        index=0,
        key="syllabus_template_choice"
    )

with col2: 
    if st.button("📋 Load Template", key="syllabus_load_template", use_container_width=True):
        # Clear existing sections
        st.session_state.course_plan['sections'] = []
        
        # Load Standard template with pre-filled structure
        for section in st.session_state.syllabus_template["Standard"]:
            if section['type'] == 'course_structure_table':
                new_section = {
                    "type": "course_structure_table",
                    "title": "Course Structure",
                    "units": [
                        {"number": 1, "title": "", "description": "", "lessons": []},
                        {"number": 2, "title": "", "description": "", "lessons": []},
                        {"number": 3, "title": "", "description": "", "lessons": []},
                        {"number": 4, "title": "", "description": "", "lessons": []}
                    ]
                }
                st.session_state.course_plan['sections'].append(new_section)
            
            elif section['type'] == 'lesson_structure_table':
                new_section = {
                    "type": "lesson_structure_table",
                    "title": "Generic Lesson Structure",
                    "stages": [
                        {"name": "WARM-UP (KHỞI ĐỘNG)", "description": ""},
                        {"name": "INTRODUCTION (GIỚI THIỆU MỤC TIÊU)", "description": ""},
                        {"name": "PRESENTATION (NỘI DUNG BÀI MỚI)", "description": ""},
                        {"name": "PRACTICE (HỌC SINH THỰC HÀNH)", "description": ""},
                        {"name": "PRODUCTION (HỌC SINH VẬN DỤNG)", "description": ""},
                        {"name": "REVIEW (ÔN TẬP)", "description": ""}
                    ]
                }
                st.session_state.course_plan['sections'].append(new_section)
            
            elif section['type'] == 'schedule':
                new_section = {
                    "type": "schedule",
                    "title": "Schedule & Lesson Information",
                    "lessons": []
                }
                st.session_state.course_plan['sections'].append(new_section)
            
            else:  # 'text' type
                new_section = {
                    "type": "text",
                    "title": section['title'],
                    "value": ""
                }
                st.session_state.course_plan['sections'].append(new_section)
        
        st.rerun()

st.divider()


# Main editing area
st.markdown("## Edit Syllabus")

# Quick add buttons - ALWAYS VISIBLE
st.markdown("### Add New Section")
col1, col2, col3, col4 = st.columns(4)
with col1:
    if st.button("➕ Add Text Section", use_container_width=True):
        st.session_state.course_plan["sections"].append({
            "type": "text", 
            "title": "New Text Section", 
            "value": ""
        })
        st.rerun()
with col2:
    if st.button("➕ Add Course Structure", use_container_width=True):
        st.session_state.course_plan["sections"].append({
            "type": "course_structure_table", 
            "title": "Course Structure", 
            "units": []
        })
        st.rerun()
with col3:
    if st.button("➕ Add Lesson Structure", use_container_width=True):
        st.session_state.course_plan["sections"].append({
            "type": "lesson_structure_table", 
            "title": "Lesson Structure", 
            "stages": []
        })
        st.rerun()
with col4:
    if st.button("➕ Add Schedule", use_container_width=True):
        st.session_state.course_plan["sections"].append({
            "type": "schedule", 
            "title": "Schedule", 
            "lessons": []
        })
        st.rerun()

st.divider()


if st.session_state.course_plan["sections"]:
    st.markdown("### Current Sections")
    for idx, section in enumerate(st.session_state.course_plan["sections"]):
        with st.expander(f"📌 {section.get('title', 'Untitled')}", expanded=True):
            col1, col2 = st.columns([4, 1])
            
            with col2:
                cols = st.columns([1,1,2])
                with cols[0]:
                    if idx > 0 and st.button("⬆️", key=f"syllabus_up_{idx}"):
                        st.session_state.course_plan["sections"][idx], st.session_state.course_plan["sections"][idx-1] = \
                            st.session_state.course_plan["sections"][idx-1], st.session_state.course_plan["sections"][idx]
                        st.rerun()
                with cols[1]:
                    if idx < len(st.session_state.course_plan["sections"])-1 and st.button("⬇️", key=f"syllabus_down_{idx}"):
                        st.session_state.course_plan["sections"][idx], st.session_state.course_plan["sections"][idx+1] = \
                            st.session_state.course_plan["sections"][idx+1], st.session_state.course_plan["sections"][idx]
                        st.rerun()
                with cols[2]:
                    if st.button("Delete", key=f"syllabus_del_{idx}"):
                        st.session_state.course_plan["sections"].pop(idx)
                        st.rerun()
            
            with col1:
                # Edit section title
                new_title = st.text_input("Title", value=section.get('title', ''), 
                                        key=f"syllabus_title_{idx}", label_visibility="collapsed")
                section['title'] = new_title
                
                # Render based on section type
                if section['type'] == 'text':
                    content = section.get('value', '')
                    if isinstance(content, list):
                        content = "\n".join(content)
                    edited_content = st.text_area(
                        "Content",
                        value=content,
                        key=f"syllabus_text_{idx}",
                        height=150,
                        placeholder="Enter your text here..."
                    )
                    section['value'] = edited_content
                
                elif section['type'] == 'course_structure_table':
                    units = section.get('units', [])
                    st.markdown("**Course Units & Lessons**")
                    
                    # Controls for adding/removing units
                    col1, col2, col3 = st.columns([2, 1, 1])
                    with col1:
                        st.markdown(f"**Total Units: {len(units)}**")
                    with col2:
                        if st.button("➕ Add Unit", key=f"syllabus_add_unit_{idx}"):
                            new_unit_num = len(units) + 1
                            units.append({
                                "number": new_unit_num,
                                "title": "",
                                "description": "",
                                "lessons": []
                            })
                            st.rerun()
                    with col3:
                        if len(units) > 0 and st.button("➖ Remove Last Unit", key=f"syllabus_remove_unit_{idx}"):
                            units.pop()
                            st.rerun()
                    
                    st.divider()
                    
                    if not units:
                        st.info("No units yet. Click 'Add Unit' to start building your course structure.")
                    
                    for u_idx, unit in enumerate(units):
                        with st.container():
                            st.markdown(f"**Unit {unit.get('number', u_idx+1)}**")
                            col_a, col_b = st.columns(2)
                            with col_a:
                                unit['title'] = st.text_input(
                                    f"Unit Title",
                                    value=unit.get('title', ''),
                                    key=f"syllabus_unit_title_{idx}_{u_idx}"
                                )
                            with col_b:
                                unit['description'] = st.text_input(
                                    "Description",
                                    value=unit.get('description', ''),
                                    key=f"syllabus_unit_desc_{idx}_{u_idx}"
                                )
                            
                            # Lessons in this unit
                            lessons = unit.get('lessons', [])
                            
                            # Add lesson button for this unit
                            if st.button("➕ Add Lesson to this Unit", key=f"syllabus_add_lesson_{idx}_{u_idx}"):
                                lessons.append({"number": "", "title": "", "focus": ""})
                                st.rerun()
                            
                            # Display lessons
                            for l_idx, lesson in enumerate(lessons):
                                cols = st.columns([1, 3, 3, 1])
                                with cols[0]:
                                    lesson['number'] = st.text_input(
                                        "No.",
                                        value=lesson.get('number', ''),
                                        key=f"syllabus_lesson_num_{idx}_{u_idx}_{l_idx}",
                                        label_visibility="collapsed",
                                        placeholder="#"
                                    )
                                with cols[1]:
                                    lesson['title'] = st.text_input(
                                        "Title",
                                        value=lesson.get('title', ''),
                                        key=f"syllabus_lesson_title_{idx}_{u_idx}_{l_idx}",
                                        label_visibility="collapsed",
                                        placeholder="Lesson title"
                                    )
                                with cols[2]:
                                    lesson['focus'] = st.text_input(
                                        "Focus",
                                        value=lesson.get('focus', ''),
                                        key=f"syllabus_lesson_focus_{idx}_{u_idx}_{l_idx}",
                                        label_visibility="collapsed",
                                        placeholder="Lesson focus"
                                    )
                                with cols[3]:
                                    if st.button("🗑️", key=f"syllabus_del_lesson_{idx}_{u_idx}_{l_idx}"):
                                        lessons.pop(l_idx)
                                        st.rerun()
                            
                            unit['lessons'] = lessons
                            st.divider()
                    
                    section['units'] = units
                
                elif section['type'] == 'lesson_structure_table':
                    stages = section.get('stages', [])
                    st.markdown("**Lesson Stages**")
                    
                    # Controls for adding/removing stages
                    col1, col2, col3 = st.columns([2, 1, 1])
                    with col1:
                        st.markdown(f"**Total Stages: {len(stages)}**")
                    with col2:
                        if st.button("➕ Add Stage", key=f"syllabus_add_stage_{idx}"):
                            stages.append({"name": "New Stage", "description": ""})
                            st.rerun()
                    with col3:
                        if len(stages) > 0 and st.button("➖ Remove Last Stage", key=f"syllabus_remove_stage_{idx}"):
                            stages.pop()
                            st.rerun()
                    
                    st.divider()
                    
                    if not stages:
                        st.info("No stages yet. Click 'Add Stage' to build your lesson structure.")
                    
                    for s_idx, stage in enumerate(stages):
                        col_name, col_desc = st.columns([1, 3])
                        with col_name:
                            stage['name'] = st.text_input(
                                "Stage Name",
                                value=stage.get('name', f'Stage {s_idx+1}'),
                                key=f"syllabus_stage_name_{idx}_{s_idx}"
                            )
                        with col_desc:
                            stage['description'] = st.text_area(
                                "Description",
                                value=stage.get('description', ''),
                                key=f"syllabus_stage_desc_{idx}_{s_idx}",
                                height=80,
                                label_visibility="collapsed",
                                placeholder="Stage description"
                            )
                    
                    section['stages'] = stages
                
                elif section['type'] == 'schedule':
                    lessons = section.get('lessons', [])
                    st.markdown("**Schedule & Lesson Information**")
                    
                    # Form to add new schedule item - NO EXPANDER
                    st.markdown("#### ➕ Add New Schedule Item")
                    col_a, col_b, col_c = st.columns(3)
                    with col_a:
                        new_code = st.text_input("Lesson Code (e.g., 1.1)", key=f"new_code_{idx}")
                    with col_b:
                        new_title = st.text_input("Lesson Title", key=f"new_title_{idx}")
                    with col_c:
                        new_date = st.text_input("Date", key=f"new_date_{idx}")
                    
                    col_d, col_e = st.columns(2)
                    with col_d:
                        new_time = st.text_input("Time", key=f"new_time_{idx}")
                    with col_e:
                        new_topic = st.text_input("Topic", key=f"new_topic_{idx}")
                    
                    new_description = st.text_area("Description", height=80, key=f"new_desc_{idx}")
                    new_objectives = st.text_area("Learning Objectives", height=80, key=f"new_obj_{idx}")
                    new_activity = st.text_input("Interactive Activity", key=f"new_act_{idx}")
                    
                    if st.button("Add to Schedule", key=f"add_to_schedule_{idx}"):
                        if new_code:
                            lessons.append({
                                "code": new_code,
                                "title": new_title,
                                "date": new_date,
                                "time": new_time,
                                "topic": new_topic,
                                "description": new_description,
                                "objectives": new_objectives,
                                "activity": new_activity
                            })
                            st.rerun()
                    
                    st.divider()
                    
                    # Display existing schedule items - NO EXPANDERS
                    if lessons:
                        st.markdown("**Current Schedule Items**")
                        sched_to_remove = []
                        for s_idx, lesson in enumerate(lessons):
                            # Simple box using columns and markdown - no expander
                            st.markdown(f"**📌 Lesson {lesson.get('code', '')}: {lesson.get('title', '')}**")
                            
                            col_info, col_del = st.columns([5, 1])
                            with col_del:
                                if st.button("🗑️ Delete", key=f"del_sched_{idx}_{s_idx}"):
                                    sched_to_remove.append(s_idx)
                            
                            with col_info:
                                st.markdown(f"📅 **Date:** {lesson.get('date', '')} {lesson.get('time', '')}")
                                st.markdown(f"📚 **Topic:** {lesson.get('topic', '')}")
                                st.markdown(f"📝 **Description:** {lesson.get('description', '')}")
                                st.markdown(f"🎯 **Objectives:** {lesson.get('objectives', '')}")
                                st.markdown(f"💡 **Activity:** {lesson.get('activity', '')}")
                            
                            st.divider()
                        
                        # Remove marked items
                        for s_idx in sorted(sched_to_remove, reverse=True):
                            lessons.pop(s_idx)
                            st.rerun()
                        
                        section['lessons'] = lessons
                    else:
                        st.info("No schedule items yet. Use the form above to add lessons.")
else:
    st.info("No sections yet. Use the buttons above to add your first section.")

# Generate Document section
st.markdown("---")
st.markdown("## Generate Document")

if st.button("Generate Document", type="primary", key="syllabus_generate", use_container_width=True):
    with st.spinner("Creating your syllabus..."):
        doc = format_syllabus_document(st.session_state.course_plan)
        
        bio = BytesIO()
        doc.save(bio)
        bio.seek(0)
        
        title = st.session_state.course_plan["metadata"]["title"].replace(" ", "_") or "Syllabus"
        filename = f"{title}_Course_Plan.docx"
        
        st.success("Document ready!")
        st.download_button(
            label="📥 Download Syllabus",
            data=bio,
            file_name=filename,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            key="syllabus_download",
            use_container_width=True
        )
        
        # Preview
        st.markdown("---")
        st.markdown("### Document Structure Preview")
        st.markdown(f"**Course:** {st.session_state.course_plan['metadata']['title']}")
        st.markdown(f"**Teacher:** {st.session_state.course_plan['metadata']['teacher']}")
        st.markdown(f"**Sections:** {len(st.session_state.course_plan['sections'])}")
        for section in st.session_state.course_plan['sections']:
            st.markdown(f"- {section['title']} (*{section['type']}*)")