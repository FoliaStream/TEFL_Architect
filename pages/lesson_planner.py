import streamlit as st
import os 
import yaml 

from src.fe.support_functions import setup_sidebar
from src.fe.styles import HIDE_SIDEBAR_NAV, TEXT_JUSTIFIED
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from io import BytesIO

################
# --- SET UP ---
################


# --- PAGE CONFIG --- 
st.set_page_config(page_title="Lesson Planner", layout="wide")

# --- STYLES ---
st.markdown(HIDE_SIDEBAR_NAV, unsafe_allow_html=True)
st.markdown(TEXT_JUSTIFIED, unsafe_allow_html=True)

# --- CONFIG ---
with open(f"{os.getcwd()}/src/be/config.yaml", "r") as config_file:
    config = yaml.safe_load(config_file)

# --- SIDEBAR & TITLE ---
selected_page = setup_sidebar(
    pages=config['pages'],
    main_page=config['main_page']
)

# Navigation on click
if selected_page == "HOME":
    st.switch_page("HOME.py")
elif selected_page == "Whiteboard":
    st.switch_page("pages/whiteboard.py")
elif selected_page == "IELTS Preparation":
    st.switch_page("pages/ielts_preparation.py")
elif selected_page == "Lesson Planner":
    st.title(selected_page)
elif selected_page == "Syllabus Planner":
    st.switch_page("pages/syllabus_planner.py")
elif selected_page == "TEFL Theory":
    st.switch_page("pages/tefl_theory.py")



# --- SESSION STATE INIT ---

# Basic info
if "lesson_plan" not in st.session_state:
    st.session_state.lesson_plan = {
        "metadata":{
            "class":"",
            "teacher":"",
            "teacher_contact":"",
            "number_students":"",
            "level_students":"",
            "date":"today",
            "start_time":None,
            "end_time":None,
            "title":"",
            "description":""
        },
        "sections": []
    }

# Standard template
if "templates" not in st.session_state:
    st.session_state.templates = {
        "Standard":[
            {"type": "list", "title": "Aims / Objectives", "value": []},
            {"type": "list", "title": "Assumptions", "value": []},
            {"type": "list", "title": "Materials", "value": []},
            {"type": "focus_table", "title": "Target Language", "value": []},
            {"type": "stages_table", "title": "Stages", "value": []},
        ]
    }

# --- PAGE FUNCTIONS --- ??to export to support functions??

def format_section_title(doc, title):

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add blue background
    p_pr = title_para._element.get_or_add_pPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), '4472C4')  # Blue
    p_pr.append(shading)
    
    # Add white text
    title_run = title_para.add_run(title)
    title_run.font.size = Pt(14)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(255, 255, 255)
    
    return title_para




def format_table(table, column_align = None, header_row = True):

    tbl = table._tbl
    tblPr = tbl.tblPr

    # Borders
    tblBorders = OxmlElement('w:tblBorders')

    border_types = {
        'top': ('single', 12),
        'bottom': ('single', 8),
        'left': ('nil', 0),
        'right': ('nil', 0),
        'insideH': ('nil', 0),
        'insideV': ('nil', 0)
    }

    for border_name, (val, sz) in border_types.items():
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        if val != 'nil':
            border.set(qn('w:sz'), str(sz))
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), '000000') # black
        tblBorders.append(border)

    existing_borders = tblPr.find(qn('w:tblBorders'))
    if existing_borders is not None:
        tblPr.remove(existing_borders)
    tblPr.append(tblBorders)

    for row in table.rows:
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcBorders = tcPr.find(qn('w:tcBorders'))
            if tcBorders is not None:
                tcPr.remove(tcBorders)

    # Text Format
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):

            if column_align and col_idx < len(column_align):
                h_align, v_align = column_align[col_idx]
            else:
                #default
                h_align, v_align = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_VERTICAL.TOP
            
            # Set alignment 
            cell.vertical_alignment = v_align
            for paragraph in cell.paragraphs:
                paragraph.alignment = h_align

            
            # Set header bold
            if header_row and row_idx == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True

    return table



def format_document(data):

    # Create doc
    doc = Document()

    # Set margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # ==== Page 1 ====

    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lesson_title = data['metadata']['title']
    header_run = header.add_run(f"{lesson_title} LESSON PLAN")
    header_run.font.size = Pt(11)
    header_run.font.bold = True
    header_run.font.color.rgb = RGBColor(100,100,100)

    # Whitespace
    doc.add_paragraph() 

    # GENERAL INFORMATION Paragraph
    general_info_title = format_section_title(doc, "GENERAL INFORMATION")
    
    # Whitespace
    doc.add_paragraph() 

    # Lesson Details
    lesson_details_table = doc.add_table(rows=4, cols=2)

    lesson_details_table = format_table(lesson_details_table,
                                        column_align=[(WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_VERTICAL.CENTER),
                                                      (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_VERTICAL.TOP)],
                                        header_row=True)
    
    lesson_details_table.cell(0,0).text = "LESSON"
    lesson_details_table.cell(0,0).paragraphs[0].runs[0].font.bold = True
    lesson_cell = lesson_details_table.cell(0,1)
    lesson_cell.text = f"Title: {data['metadata'].get('title','')}"
    lesson_cell.paragraphs[0].add_run(f"\nDescription: {data['metadata'].get('description','')}")

    lesson_details_table.cell(1,0).text = "DATE & TIME"
    lesson_details_table.cell(1,0).paragraphs[0].runs[0].font.bold = True
    datetime_cell = lesson_details_table.cell(1,1)
    datetime_cell.text = f"Date: {data['metadata'].get('date','')}"
    datetime_cell.paragraphs[0].add_run(f"\nStart Time: {data['metadata'].get('start_time', '')}")
    datetime_cell.paragraphs[0].add_run(f"\nEnd Time: {data['metadata'].get('end_time', '')}")

    lesson_details_table.cell(2,0).text = "TEACHER CONTACT"
    lesson_details_table.cell(2,0).paragraphs[0].runs[0].font.bold = True
    contact_cell = lesson_details_table.cell(2,1)
    contact_cell.text = f"Name: {data['metadata'].get('teacher', '')}"
    contact_cell.paragraphs[0].add_run(f"\nContact: {data['metadata'].get('teacher_contact','')}")

    lesson_details_table.cell(3,0).text = "STUDENTS INFORMATION"
    lesson_details_table.cell(3,0).paragraphs[0].runs[0].font.bold = True
    students_cell = lesson_details_table.cell(3,1)
    students_cell.text = f"Number of Students: {data['metadata'].get('number_students', '')}"
    students_cell.paragraphs[0].add_run(f"\nLevel of Students: {data['metadata'].get('level_students','')}")


    # Whitespace
    doc.add_paragraph() 

    # Find and render each section 
    for section in data['sections']:
        section_title_text = format_section_title(doc, section['title'])

        if section['type'] == 'stages_table':
            
            stages = section.get('stages',[])
            if stages:
                stages_table = doc.add_table(rows=1, cols=3)
                stages_table.style = "Table Grid"
                    
                stages_table = format_table(stages_table,
                                            column_align=[(WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_VERTICAL.CENTER),
                                                          (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_VERTICAL.CENTER),
                                                          (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_VERTICAL.TOP)],
                                                          header_row=True)
                headers = stages_table.rows[0].cells
                headers[0].text = "STAGE"
                headers[1].text = "DURATION"
                headers[2].text = "DESCRIPTION"

                for stage in stages:
                    if stage.get('name') or stage.get('description'):
                        row = stages_table.add_row().cells
                        row[0].text = stage.get('name', '')
                        row[1].text = stage.get('duration', '')
                        row[2].text = stage.get('description', '')
            else:
                doc.add_paragraph("(No stages added)").italic = True
        
        elif section['type'] == 'list': 
            section_table = doc.add_table(rows=1,cols=1)
            section_table.style = "Table Grid"

            section_table = format_table(section_table, 
                                         column_align=[(WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_VERTICAL.TOP)],
                                         header_row=False)
            
            content_cell = section_table.cell(0,0)
            content_cell.text = ""

            section_items = section.get('value', [])
            for i, item in enumerate(section_items):
                if i>0:
                    content_cell.paragraphs[0].add_run("\n")
                content_cell.paragraphs[0].add_run(f"• {item}")

        elif section['type'] in ['number', 'text']:
            section_value = section.get('value','')
            section_table = doc.add_table(rows=1,cols=1)
            section_table.style = "Table Grid"

            section_table = format_table(section_table, 
                                         column_align=[(WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_VERTICAL.TOP)],
                                         header_row=False)
            
            content_cell = section_table.cell(0,0)
            content_cell.paragraphs[0].add_run(str(section_value))

        elif section['type'] == 'focus_table':

            
            focus_table = doc.add_table(rows=3, cols=2)
            focus_table.style = "Table Grid"
            focus_table = format_table(focus_table,
                                        column_align=[(WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_VERTICAL.CENTER),
                                                        (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_VERTICAL.TOP)],
                                        header_row=False)
        
            focus_table.cell(0,0).text = "VOCABULARY"
            focus_table.cell(0,0).paragraphs[0].runs[0].font.bold = True
            vocabulary_cell = focus_table.cell(0,1)
            vocabulary_cell.text = ""
            vocabulary_items = section.get('vocabulary',[])
            if vocabulary_items:
                vocabulary_cell.paragraphs[0].add_run(", ".join(vocabulary_items))

            focus_table.cell(1,0).text = "GRAMMAR"
            focus_table.cell(1,0).paragraphs[0].runs[0].font.bold = True
            grammar_cell = focus_table.cell(1,1)
            grammar_cell.text = ""
            grammar_items = section.get('grammar', [])
            if grammar_items:
                grammar_cell.paragraphs[0].add_run(", ".join(grammar_items))
            
            focus_table.cell(2,0).text = "FUNCTIONAL LANGUAGE"
            focus_table.cell(2,0).paragraphs[0].runs[0].font.bold = True
            func_cell = focus_table.cell(2,1)
            func_cell.text = ""
            func_text = section.get('functional_language', '')
            if func_text:
                func_cell.paragraphs[0].add_run(func_text)


    return doc


# --- PAGE ---

# Input Basic Information

st.divider()

col1, col2 = st.columns(2)
with col1:
    st.markdown("### Lesson Information")
    st.session_state.lesson_plan["metadata"]["title"] = st.text_input(
        "Lesson Title", value=st.session_state.lesson_plan["metadata"]["title"]
    )
    st.session_state.lesson_plan["metadata"]["number_students"] = st.text_input(
        "Number of Students", value=st.session_state.lesson_plan["metadata"]["number_students"]
    )
    st.session_state.lesson_plan["metadata"]["level_students"] = st.text_input(
        "Level of Students", value=st.session_state.lesson_plan["metadata"]["level_students"]
    )
    st.session_state.lesson_plan["metadata"]["description"] = st.text_area(
        "Description", value=st.session_state.lesson_plan["metadata"]["description"], height=125
    )

with col2:
    st.markdown("### Schedule & More")
    st.session_state.lesson_plan["metadata"]["date"] = st.date_input(
        "Date", value=st.session_state.lesson_plan["metadata"]["date"]
    )
    st.session_state.lesson_plan["metadata"]["start_time"] = st.time_input(
        "Start Time", value=st.session_state.lesson_plan["metadata"]["start_time"]
    )
    st.session_state.lesson_plan["metadata"]["end_time"] = st.time_input(
        "End Time", value=st.session_state.lesson_plan["metadata"]["end_time"]
    )
    st.session_state.lesson_plan["metadata"]["teacher"] = st.text_input(
        "Teacher Name", value=st.session_state.lesson_plan["metadata"]["teacher"]
    )
    st.session_state.lesson_plan["metadata"]["teacher_contact"] = st.text_input(
        "Teacher Contact", value=st.session_state.lesson_plan["metadata"]["teacher_contact"]
    )

st.divider()


# Load Std Template
col1, col2 = st.columns([3,1])
with col1: 
    template_choice = st.selectbox(
        "Load Template",
        options=list(st.session_state.templates.keys()),
        index=0
    )

with col2: 
    if st.button("📋 Load Template", use_container_width=True):
        # Clear existing sections
        st.session_state.lesson_plan['sections'] = []
        
        # Load Standard template with empty values
        for section in st.session_state.templates["Standard"]:
            if section['type'] == 'stages_table':
                # Empty stages table
                new_section = section.copy()
                new_section['stages'] = []  # Empty list
                st.session_state.lesson_plan['sections'].append(new_section)
            
            elif section['type'] == 'focus_table':
                # Empty focus table
                new_section = section.copy()
                new_section['vocabulary'] = []
                new_section['grammar'] = []
                new_section['functional_language'] = ''
                st.session_state.lesson_plan['sections'].append(new_section)
            
            else:  # 'list' and other types
                new_section = section.copy()
                if new_section['type'] == 'list':
                    new_section['value'] = []  # Empty list
                st.session_state.lesson_plan['sections'].append(new_section)
        
        st.rerun()

st.divider()



tab1, tab2, tab3 = st.tabs(["Edit Lesson Plan", "Add Sections", "Generate Document"])

with tab1:
    if not st.session_state.lesson_plan["sections"]:
        st.info("Load the Standard template or add sections manually")
    else:
        for idx, section in enumerate(st.session_state.lesson_plan["sections"]):
            with st.expander(f"📌 {section.get('title', 'Untitled')}", expanded=True):
                col1, col2 = st.columns([4, 1])
                
                with col2:
                    cols = st.columns([1,1,2])
                    with cols[0]:
                        if idx > 0 and st.button("⬆️", key=f"up_{idx}"):
                            st.session_state.lesson_plan["sections"][idx], st.session_state.lesson_plan["sections"][idx-1] = \
                                st.session_state.lesson_plan["sections"][idx-1], st.session_state.lesson_plan["sections"][idx]
                            st.rerun()
                    with cols[1]:
                        if idx < len(st.session_state.lesson_plan["sections"])-1 and st.button("⬇️", key=f"down_{idx}"):
                            st.session_state.lesson_plan["sections"][idx], st.session_state.lesson_plan["sections"][idx+1] = \
                                st.session_state.lesson_plan["sections"][idx+1], st.session_state.lesson_plan["sections"][idx]
                            st.rerun()
                    with cols[2]:
                        if st.button("Delete", key=f"del_{idx}"):
                            st.session_state.lesson_plan["sections"].pop(idx)
                            st.rerun()
                
                with col1:
                    # Edit section title
                    new_title = st.text_input("Title", value=section.get('title', ''), 
                                            key=f"title_{idx}", label_visibility="collapsed")
                    section['title'] = new_title
                    
                    # Render based on section type
                    if section['type'] == 'list':
                        items = section.get('value', [])
                        item_text = st.text_area(
                            "Items (one per line)", 
                            value="\n".join(items),
                            key=f"val_{idx}", height=100
                        )
                        section['value'] = [line.strip() for line in item_text.split('\n') if line.strip()]
                    
                    elif section['type'] == 'stages_table':
                        stages = section.get('stages', [])
                        st.markdown("**Stage | Duration | Description**")
                        
                        for s_idx, stage in enumerate(stages):
                            cols = st.columns([2, 1, 3, 1])
                            with cols[0]:
                                stage['name'] = st.text_input("Stage", value=stage.get('name', ''),
                                                            key=f"stage_name_{idx}_{s_idx}")
                            with cols[1]:
                                stage['duration'] = st.text_input("Duration", value=stage.get('duration', ''),
                                                                key=f"stage_dur_{idx}_{s_idx}")
                            with cols[2]:
                                stage['description'] = st.text_area("Description", value=stage.get('description', ''),
                                                                    key=f"stage_desc_{idx}_{s_idx}")
                            with cols[3]:
                                if st.button("🗑️", key=f"del_stage_{idx}_{s_idx}"):
                                    stages.pop(s_idx)
                                    st.rerun()
                        
                        if st.button("➕ Add Stage", key=f"add_stage_{idx}"):
                            stages.append({"name": "", "duration": "", "description": ""})
                            st.rerun()
                        
                        section['stages'] = stages
                    
                    elif section['type'] == 'focus_table':
                        st.markdown("**Vocabulary**")
                        vocab_items = section.get('vocabulary', [])
                        vocab_text = st.text_area(
                            "Vocabulary (comma-separated)", 
                            value=", ".join(vocab_items),
                            key=f"vocab_{idx}", height=80
                        )
                        section['vocabulary'] = [v.strip() for v in vocab_text.split(',') if v.strip()]
                        
                        st.markdown("**Grammar**")
                        grammar_items = section.get('grammar', [])
                        grammar_text = st.text_area(
                            "Grammar (comma-separated)", 
                            value=", ".join(grammar_items),
                            key=f"grammar_{idx}", height=80
                        )
                        section['grammar'] = [g.strip() for g in grammar_text.split(',') if g.strip()]
                        
                        st.markdown("**Functional Language**")
                        func_text = section.get('functional_language', '')
                        section['functional_language'] = st.text_area(
                            "Functional Language",
                            value=func_text,
                            key=f"func_{idx}", height=100
                        )

with tab2:
    st.markdown("### Add New Section")
    
    col1, col2 = st.columns(2)
    with col1:
        new_type = st.selectbox(
            "Section Type",
            ["list", "stages_table", "focus_table"],
            format_func=lambda x: {
                "list": "Simple List",
                "stages_table": "Stages Table",
                "focus_table": "Target Language (Vocab/Grammar/Func)"
            }[x]
        )
    with col2:
        if st.button("Add Section", type="primary", use_container_width=True):
            new_section = {"type": new_type, "title": new_type.replace("_", " ").title()}
            
            if new_type == "list":
                new_section["value"] = []
            
            elif new_type == "stages_table":
                new_section["stages"] = []
            
            elif new_type == "focus_table":
                new_section["vocabulary"] = []
                new_section["grammar"] = []
                new_section["functional_language"] = ""
            
            st.session_state.lesson_plan["sections"].append(new_section)
            st.rerun()

with tab3:
    st.markdown("### Generate Lesson Plan Document")
    
    if st.button("Generate Document", type="primary", use_container_width=True):
        with st.spinner("Creating your lesson plan..."):
            doc = format_document(st.session_state.lesson_plan)
            
            bio = BytesIO()
            doc.save(bio)
            bio.seek(0)
            
            title = st.session_state.lesson_plan["metadata"]["title"].replace(" ", "_") or "Lesson"
            filename = f"{title}_Lesson_Plan.docx"
            
            st.success("Document ready!")
            st.download_button(
                label="📥 Download Lesson Plan",
                data=bio,
                file_name=filename,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )
            
            # Preview
            st.markdown("---")
            st.markdown("### Document Structure Preview")
            st.markdown(f"**Title:** {st.session_state.lesson_plan['metadata']['title']}")
            st.markdown(f"**Teacher:** {st.session_state.lesson_plan['metadata']['teacher']}")
            st.markdown(f"**Sections:** {len(st.session_state.lesson_plan['sections'])}")
            for section in st.session_state.lesson_plan['sections']:
                st.markdown(f"- {section['title']} (*{section['type']}*)")



